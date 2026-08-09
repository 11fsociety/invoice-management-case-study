"""
Generate 9 edge-case invoice PDFs (INV-0021 through INV-0029) for the Zamp ASA
case study. Reuses the layout helpers from generate_invoices.py. Idempotent -
skips PDFs that already exist. Also appends 4 new PO rows to po_master.xlsx
and po_master.csv.
"""

from __future__ import annotations
import os
import sys
import csv
from decimal import Decimal, ROUND_HALF_UP
from datetime import date
from pathlib import Path

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Table, TableStyle

# Reuse layout helpers from generate_invoices.py
HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
from generate_invoices import (  # type: ignore
    money,
    compute_totals,
    draw_left_header,
    draw_right_header,
    draw_center,
    draw_twocol,
    draw_boxed,
)

OUT_DIR = Path(r"D:\codezzz\Claude\zamp-asa\sample-invoices")
OUT_DIR.mkdir(parents=True, exist_ok=True)


# Edge invoice entries. Each tuple matches the layout of generate_invoices.py.
# (num, vendor, vendor_id, po_number, po_date, invoice_date, layout, line_items, tax_style, currency, po_ref_style)
EDGE_INVOICES = [
    (21, "Bharat Auto Spares", "GST-27AABCB1234M1Z5", "PO-2001", date(2026, 7, 5), date(2026, 7, 10), "left",
        [("Consulting engagement - Milestone 1", 1, 60000)], "none", "INR", "PO#"),
    (22, "Bharat Auto Spares", "GST-27AABCB1234M1Z5", "PO-2001", date(2026, 7, 5), date(2026, 7, 15), "right",
        [("Consulting engagement - Milestone 2", 1, 40000)], "none", "INR", "PO#"),
    (23, "Bharat Auto Spares", "GST-27AABCB1234M1Z5", "PO-2001", date(2026, 7, 5), date(2026, 7, 20), "center",
        [("Consulting engagement - Milestone 3", 1, 20000)], "none", "INR", "PO#"),
    # INV-0024: wrong printed GST id -> vendor identity mismatch
    (24, "Kumar Traders", "GST-CONFLICT-999", "PO-2002", date(2026, 7, 8), date(2026, 7, 14), "boxed",
        [("Q3 consulting", 40, 2500)], "none", "INR", "Purchase Order:"),
    # INV-0025: "Kushaq" alias should resolve against PO "New Kushaq"
    (25, "Volkswagen Components India", "GST-06AAACV5566H1ZP", "PO-2003", date(2026, 7, 1), date(2026, 7, 10), "twocol",
        [("Kushaq", 2, 800000)], "none", "INR", "Ref:"),
    # INV-0026: exact duplicate of INV-0001 (same content). We override the
    # printed invoice number to "INV-0001" so the LLM extracts that value.
    (26, "Bharat Auto Spares", "GST-27AABCB1234M1Z5", "PO-1001", date(2026, 7, 15), date(2026, 7, 22), "left",
        [("Brake pad set", 20, 450), ("Air filter", 10, 320), ("Oil seal", 15, 180)], "single", "INR", "PO#"),
    # INV-0027: credit note (special layout)
    (27, "Bharat Auto Spares", "GST-27AABCB1234M1Z5", "PO-1001", date(2026, 7, 15), date(2026, 7, 25), "credit_note",
        [("Credit for damaged brake pad returns", 1, 10000)], "none", "INR", "PO#"),
    # INV-0028: over-tolerance
    (28, "Local Cartons Co", "GST-27AABLC2233N1ZY", "PO-2004", date(2026, 7, 10), date(2026, 7, 20), "boxed",
        [("Carton stock", 25, 2104)], "none", "INR", "PO#"),
    # INV-0029: NO PO reference anywhere
    (29, "Local Cartons Co", "GST-27AABLC2233N1ZY", None, None, date(2026, 7, 22), "no_po",
        [("Miscellaneous supplies", 1, 15000)], "none", "INR", None),
]


def draw_credit_note(c, inv, styles):
    """A credit-note layout - similar to 'left' but with a big CREDIT NOTE title
    and a prominent 'Reference Invoice' row. Total shown as negative."""
    W, H = A4
    num, vendor, vid, po, po_date, inv_date, _, items, tax_style, cur, po_ref = inv
    subtotal, tax, total = compute_totals(items, tax_style)
    total_negative = -total

    c.setFont("Helvetica-Bold", 20)
    c.drawString(20 * mm, H - 25 * mm, vendor)
    c.setFont("Helvetica", 9)
    c.drawString(20 * mm, H - 32 * mm, f"Vendor ID: {vid}")
    c.drawString(20 * mm, H - 37 * mm, "Registered: Industrial Estate, Sector 7")
    c.setStrokeColor(colors.HexColor("#aa0000"))
    c.setLineWidth(2)
    c.line(20 * mm, H - 42 * mm, W - 20 * mm, H - 42 * mm)
    c.setLineWidth(1)
    c.setFillColor(colors.HexColor("#aa0000"))
    c.setFont("Helvetica-Bold", 22)
    c.drawString(20 * mm, H - 55 * mm, "CREDIT NOTE / DEBIT NOTE")
    c.setFillColor(colors.black)
    c.setFont("Helvetica", 10)
    c.drawString(20 * mm, H - 65 * mm, f"Credit Note No: INV-{num:04d}")
    c.drawString(20 * mm, H - 71 * mm, f"Date: {inv_date.strftime('%d %b %Y')}")
    c.drawString(120 * mm, H - 65 * mm, f"{po_ref} {po}")
    c.drawString(120 * mm, H - 71 * mm, f"PO Date: {po_date.strftime('%d %b %Y')}")

    c.setFont("Helvetica-Bold", 12)
    c.drawString(20 * mm, H - 82 * mm, "Reference Invoice: INV-0001")
    c.setFont("Helvetica", 9)
    c.drawString(20 * mm, H - 88 * mm,
        "This credit note adjusts the referenced invoice for returned/damaged goods.")

    y = H - 100 * mm
    data = [["#", "Description", "Qty", "Rate", "Amount"]]
    for i, (desc, qty, rate) in enumerate(items, 1):
        amt = Decimal(str(qty)) * Decimal(str(rate))
        data.append([str(i), desc, str(qty), money(rate, cur), money(amt, cur)])
    t = Table(data, colWidths=[10 * mm, 90 * mm, 15 * mm, 30 * mm, 30 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ffe5e5")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (2, 1), (4, -1), "RIGHT"),
    ]))
    w, h = t.wrapOn(c, W, H)
    t.drawOn(c, 20 * mm, y - h)
    y = y - h - 10 * mm
    c.setFont("Helvetica", 10)
    c.drawRightString(W - 20 * mm, y, f"Subtotal: {money(subtotal, cur)}")
    y -= 6 * mm
    c.setFont("Helvetica-Bold", 14)
    c.setFillColor(colors.HexColor("#aa0000"))
    c.drawRightString(W - 20 * mm, y, f"Credit Total: -{money(total, cur)}")
    c.setFillColor(colors.black)
    c.setFont("Helvetica-Oblique", 8)
    c.drawString(20 * mm, 20 * mm,
        "Please deduct this amount from the referenced invoice INV-0001.")


def draw_no_po(c, inv, styles):
    """A minimal invoice layout that never prints any PO reference."""
    W, H = A4
    num, vendor, vid, po, po_date, inv_date, _, items, tax_style, cur, po_ref = inv
    subtotal, tax, total = compute_totals(items, tax_style)

    c.setFont("Helvetica-Bold", 18)
    c.drawString(20 * mm, H - 25 * mm, vendor)
    c.setFont("Helvetica", 9)
    c.drawString(20 * mm, H - 32 * mm, f"Vendor ID: {vid}")
    c.drawString(20 * mm, H - 37 * mm, "Contact: billing@example.com")
    c.line(20 * mm, H - 42 * mm, W - 20 * mm, H - 42 * mm)

    c.setFont("Helvetica-Bold", 14)
    c.drawString(20 * mm, H - 52 * mm, "INVOICE")
    c.setFont("Helvetica", 10)
    c.drawString(20 * mm, H - 60 * mm, f"Invoice No: INV-{num:04d}")
    c.drawString(20 * mm, H - 66 * mm, f"Date: {inv_date.strftime('%d %b %Y')}")

    y = H - 80 * mm
    data = [["Description", "Qty", "Rate", "Amount"]]
    for desc, qty, rate in items:
        amt = Decimal(str(qty)) * Decimal(str(rate))
        data.append([desc, str(qty), money(rate, cur), money(amt, cur)])
    t = Table(data, colWidths=[110 * mm, 15 * mm, 25 * mm, 25 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (1, 1), (3, -1), "RIGHT"),
    ]))
    w, h = t.wrapOn(c, W, H)
    t.drawOn(c, 20 * mm, y - h)
    y = y - h - 8 * mm
    c.setFont("Helvetica", 10)
    c.drawRightString(W - 20 * mm, y, f"Subtotal: {money(subtotal, cur)}")
    y -= 5 * mm
    c.setFont("Helvetica-Bold", 12)
    c.drawRightString(W - 20 * mm, y, f"Total Due: {money(total, cur)}")
    c.setFont("Helvetica-Oblique", 8)
    c.drawString(20 * mm, 20 * mm, "Payment terms: Net 30.")


def draw_left_forced_invno(c, inv, styles, forced_invno):
    """Same as draw_left_header but overrides the printed invoice number."""
    W, H = A4
    num, vendor, vid, po, po_date, inv_date, _, items, tax_style, cur, po_ref = inv
    subtotal, tax, total = compute_totals(items, tax_style)
    c.setFont("Helvetica-Bold", 20)
    c.drawString(20 * mm, H - 25 * mm, vendor)
    c.setFont("Helvetica", 9)
    c.drawString(20 * mm, H - 32 * mm, f"Vendor ID: {vid}")
    c.drawString(20 * mm, H - 37 * mm, "123 Industrial Estate, Sector 7")
    c.drawString(20 * mm, H - 42 * mm, "Contact: accounts@example.com | +91-22-5555-1200")
    c.setStrokeColor(colors.HexColor("#333333"))
    c.line(20 * mm, H - 46 * mm, W - 20 * mm, H - 46 * mm)
    c.setFont("Helvetica-Bold", 14)
    c.drawString(20 * mm, H - 56 * mm, "TAX INVOICE")
    c.setFont("Helvetica", 10)
    c.drawString(20 * mm, H - 64 * mm, f"Invoice No: {forced_invno}")
    c.drawString(20 * mm, H - 70 * mm, f"Invoice Date: {inv_date.strftime('%d %b %Y')}")
    c.drawString(120 * mm, H - 64 * mm, f"{po_ref} {po}")
    c.drawString(120 * mm, H - 70 * mm, f"PO Date: {po_date.strftime('%d %b %Y')}")

    y = H - 85 * mm
    data = [["#", "Description", "Qty", "Rate", "Amount"]]
    for i, (desc, qty, rate) in enumerate(items, 1):
        amt = Decimal(str(qty)) * Decimal(str(rate))
        data.append([str(i), desc, str(qty), money(rate, cur), money(amt, cur)])
    t = Table(data, colWidths=[10 * mm, 90 * mm, 15 * mm, 30 * mm, 30 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (2, 1), (4, -1), "RIGHT"),
    ]))
    w, h = t.wrapOn(c, W, H)
    t.drawOn(c, 20 * mm, y - h)
    y = y - h - 8 * mm
    c.setFont("Helvetica", 10)
    c.drawRightString(W - 20 * mm, y, f"Subtotal: {money(subtotal, cur)}")
    y -= 5 * mm
    if tax_style == "single":
        c.drawRightString(W - 20 * mm, y, f"GST @ 18%: {money(tax, cur)}")
        y -= 5 * mm
    elif tax_style == "embedded":
        c.drawRightString(W - 20 * mm, y, f"(includes GST @ 18%: {money(tax, cur)})")
        y -= 5 * mm
    c.setFont("Helvetica-Bold", 12)
    c.drawRightString(W - 20 * mm, y, f"Total: {money(total, cur)}")
    c.setFont("Helvetica-Oblique", 8)
    c.drawString(20 * mm, 20 * mm, "Payment terms: Net 30. All amounts in " + cur + ".")


LAYOUT_MAP = {
    "left": draw_left_header,
    "right": draw_right_header,
    "center": draw_center,
    "twocol": draw_twocol,
    "boxed": draw_boxed,
    "credit_note": draw_credit_note,
    "no_po": draw_no_po,
}


NEW_POS = [
    ("PO-2001", "2026-07-05", "Bharat Auto Spares", "GST-27AABCB1234M1Z5", 100000, "INR",
        "Consulting engagement - Milestone 1; Consulting engagement - Milestone 2"),
    ("PO-2002", "2026-07-08", "Kumar Traders", "PAN-AAAPK1234C", 100000, "INR",
        "Q3 consulting"),
    ("PO-2003", "2026-07-01", "Volkswagen Components India", "GST-06AAACV5566H1ZP", 1600000, "INR",
        "New Kushaq"),
    ("PO-2004", "2026-07-10", "Local Cartons Co", "GST-27AABLC2233N1ZY", 50000, "INR",
        "Carton stock replenishment"),
]


def main():
    styles = getSampleStyleSheet()

    for inv in EDGE_INVOICES:
        num = inv[0]
        layout = inv[6]
        path = OUT_DIR / f"INV-{num:04d}.pdf"
        if path.exists():
            print(f"[skip] {path.name} exists")
            continue

        c = canvas.Canvas(str(path), pagesize=A4)
        if num == 26:
            # Force invoice number to say INV-0001
            draw_left_forced_invno(c, inv, styles, "INV-0001")
        else:
            LAYOUT_MAP[layout](c, inv, styles)
        c.showPage()
        c.save()
        print(f"[ok] {path.name}")

    # Update PO master xlsx (append new POs if missing).
    try:
        from openpyxl import load_workbook, Workbook
    except ImportError:
        os.system("pip install openpyxl > NUL 2>&1")
        from openpyxl import load_workbook, Workbook

    xlsx_path = OUT_DIR / "po_master.xlsx"
    if xlsx_path.exists():
        wb = load_workbook(str(xlsx_path))
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "PO Master"
        ws.append(["po_number", "po_date", "vendor_name", "vendor_id", "po_amount", "currency", "line_item_summary"])

    existing = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            existing.add(str(row[0]).strip())

    for po in NEW_POS:
        if po[0] in existing:
            print(f"[skip] xlsx already has {po[0]}")
            continue
        ws.append(list(po))
        print(f"[ok] xlsx appended {po[0]}")
    wb.save(str(xlsx_path))

    # Also mirror to po_master.csv.
    csv_path = OUT_DIR / "po_master.csv"
    if csv_path.exists():
        with open(csv_path, "r", encoding="utf-8") as f:
            existing_csv = set()
            reader = csv.reader(f)
            next(reader, None)
            for row in reader:
                if row and row[0]:
                    existing_csv.add(row[0].strip())
        with open(csv_path, "a", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            for po in NEW_POS:
                if po[0] in existing_csv:
                    print(f"[skip] csv already has {po[0]}")
                    continue
                writer.writerow(po)
                print(f"[ok] csv appended {po[0]}")
    else:
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["po_number", "po_date", "vendor_name", "vendor_id",
                             "po_amount", "currency", "line_item_summary"])
            for po in NEW_POS:
                writer.writerow(po)


if __name__ == "__main__":
    main()
