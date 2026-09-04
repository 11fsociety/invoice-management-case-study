"""Export po_master.xlsx to CSV for embedding in the PRD."""
from openpyxl import load_workbook
from pathlib import Path
import csv

src = Path(r"D:\codezzz\Claude\invoice-management-case-study\sample-invoices\po_master.xlsx")
dst = Path(r"D:\codezzz\Claude\invoice-management-case-study\sample-invoices\po_master.csv")

wb = load_workbook(src)
ws = wb.active
with open(dst, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    for row in ws.iter_rows(values_only=True):
        w.writerow(row)
print("[ok]", dst)
